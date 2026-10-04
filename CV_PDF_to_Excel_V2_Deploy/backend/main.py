from __future__ import annotations

import io
import json
import os
import re
import zipfile
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any

import fitz
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter

APP_PASSWORD = os.getenv("APP_PASSWORD", "").strip()
FRONTEND_URLS = [
    x.strip()
    for x in os.getenv("FRONTEND_URL", "http://localhost:3000").split(",")
    if x.strip()
]

BASE_DIR = Path(__file__).resolve().parent
TEMPLATE = BASE_DIR.parent / "template" / "OLA_Annexure_XXI_template.xlsx"

app = FastAPI(
    title="CV PDF to Excel Converter",
    version="2.0.0",
    description="Extract Tamil Nadu Legal Metrology verification certificates into the supplied Annexure-XXI Excel template.",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=FRONTEND_URLS,
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


def clean(value: str | None) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def number_value(text: str, label_pattern: str) -> float | None:
    match = re.search(
        label_pattern + r"\s*[:\-]?\s*([0-9][0-9,]*(?:\.[0-9]+)?)",
        text,
        re.IGNORECASE,
    )
    if not match:
        return None
    return float(match.group(1).replace(",", ""))


def extract_customer(text: str) -> str:
    # Primary: stop immediately before ", Locality".
    match = re.search(
        r"belonging\s+to\s+M/S\s*(.*?)(?=,\s*Locality\b)",
        text,
        re.IGNORECASE | re.DOTALL,
    )
    if match:
        return clean(match.group(1))

    # Fallback: stop at the first comma after "belonging to M/S".
    match = re.search(
        r"belonging\s+to\s+M/S\s*([^,\n]+)",
        text,
        re.IGNORECASE,
    )
    return clean(match.group(1)) if match else ""


def extract_class(text: str) -> str:
    patterns = [
        r"(?:^|\n|\s)(I{1,3})\s+Make\s*:",
        r"-\s*-\s*(I{1,3})\s+Make\s*:",
    ]
    for pattern in patterns:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            return match.group(1).upper()
    return ""


def extract_machine(text: str) -> str:
    match = re.search(
        r"Machine\s*No\.\s*:\s*([A-Z0-9][A-Z0-9\-]*)",
        text,
        re.IGNORECASE,
    )
    if not match:
        return ""

    machine = match.group(1).replace(" ", "")
    tail = text[match.end():]

    # The supplied certificates print a numeric continuation on the next
    # token (e.g. DC810213503 + 90), while the supplied Excel combines it.
    continuation = re.match(
        r"\s*(\d{1,3})(?=\s*(?:Stamping\s+Fee|50%\s+Additional|OD\s*:))",
        tail,
        re.IGNORECASE,
    )
    if continuation:
        machine += continuation.group(1)

    return machine


def extract_capacity(text: str) -> str:
    match = re.search(
        r"\bMax\s*:\s*([0-9]+(?:\.[0-9]+)?)\s*([A-Za-z]+)",
        text,
        re.IGNORECASE,
    )
    return f"{match.group(1)} {match.group(2)}" if match else ""


def parse_certificate(text: str) -> dict[str, Any]:
    text = text.replace("\u00ad", "")

    cv = re.search(
        r"CV\s*No\s*:\s*([A-Z0-9/.-]+)",
        text,
        re.IGNORECASE,
    )
    dt = re.search(
        r"\bDate\s*:\s*(\d{1,2}[-/]\d{1,2}[-/]\d{4})",
        text,
        re.IGNORECASE,
    )
    make = re.search(
        r"Make\s*:\s*([A-Za-z0-9 ._-]+)",
        text,
        re.IGNORECASE,
    )

    additional_present = bool(
        re.search(r"50%\s*Additional\s*Fee", text, re.IGNORECASE)
    )
    additional = (
        number_value(text, r"50%\s*Additional\s*Fee")
        if additional_present
        else None
    )

    od = number_value(text, r"OD\s*:")

    record = {
        "cv_number": cv.group(1).strip() if cv else "",
        "cv_date": dt.group(1).strip() if dt else "",
        "machine_number": extract_machine(text),
        "capacity": extract_capacity(text),
        "class": extract_class(text),
        "manufacturer": clean(make.group(1)) if make else "",
        "customer_name": extract_customer(text),
        "location": "HOSUR",
        "actual_stamping_fee": number_value(text, r"Stamping\s*Fee"),
        "site_stamping_fee": additional,
        "od_charges": od,
        "has_site_fee_label": additional_present,
    }

    warnings: list[str] = []
    required = {
        "CV Number": record["cv_number"],
        "CV Date": record["cv_date"],
        "Machine Number": record["machine_number"],
        "Capacity": record["capacity"],
        "Class": record["class"],
        "Customer Name": record["customer_name"],
        "Stamping Fee": record["actual_stamping_fee"],
    }
    for label, value in required.items():
        if value in ("", None):
            warnings.append(f"Missing {label}")

    record["warnings"] = warnings
    record["review_required"] = bool(warnings)
    return record


def split_certificates(text: str) -> list[str]:
    matches = list(re.finditer(r"(?=CV\s*No\s*:)", text, re.IGNORECASE))
    chunks: list[str] = []

    for i, match in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        chunk = text[match.start():end]
        if re.search(r"Machine\s*No\.", chunk, re.IGNORECASE):
            chunks.append(chunk)

    return chunks


def read_pdf(upload: bytes) -> list[str]:
    doc = fitz.open(stream=upload, filetype="pdf")
    try:
        text = "\n".join(page.get_text("text") for page in doc)
    finally:
        doc.close()

    return split_certificates(text)


def extract_file(data: bytes, filename: str) -> dict[str, Any]:
    chunks = read_pdf(data)
    certificates = []

    for index, chunk in enumerate(chunks, start=1):
        record = parse_certificate(chunk)
        record["_source_pdf"] = filename
        record["_certificate_index"] = index
        certificates.append(record)

    return {
        "filename": filename,
        "certificate_count": len(certificates),
        "certificates": certificates,
    }


def check_password(request: Request) -> None:
    if APP_PASSWORD and request.headers.get("x-app-password", "") != APP_PASSWORD:
        raise HTTPException(status_code=401, detail="Invalid app password.")


def parse_date(value: Any) -> datetime | str:
    if value in ("", None):
        return ""
    if isinstance(value, datetime):
        return value

    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(value).strip(), fmt)
        except ValueError:
            pass

    return str(value).strip()


def filename_stem(name: str) -> str:
    stem = Path(name or "output").stem
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", stem).strip("._-")
    return stem or "output"


def copy_row_format(ws, source_row: int, target_row: int) -> None:
    if target_row != source_row:
        for col in range(1, ws.max_column + 1):
            source = ws.cell(source_row, col)
            target = ws.cell(target_row, col)
            if source.has_style:
                target._style = copy(source._style)
            if source.number_format:
                target.number_format = source.number_format
            if source.alignment:
                target.alignment = copy(source.alignment)
            if source.protection:
                target.protection = copy(source.protection)
            if source.font:
                target.font = copy(source.font)
            if source.fill:
                target.fill = copy(source.fill)
            if source.border:
                target.border = copy(source.border)

        if source_row in ws.row_dimensions:
            ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def shift_footer_merges_for_extra_rows(ws, extra_rows: int) -> None:
    if extra_rows <= 0:
        return

    footer_merges = [
        str(rng)
        for rng in ws.merged_cells.ranges
        if rng.min_row >= 37
    ]

    for rng in footer_merges:
        ws.unmerge_cells(rng)

    ws.insert_rows(37, extra_rows)

    for rng in footer_merges:
        match = re.match(
            r"([A-Z]+)(\d+):([A-Z]+)(\d+)",
            rng,
        )
        if not match:
            continue

        min_col, min_row, max_col, max_row = match.groups()
        min_row = int(min_row) + extra_rows
        max_row = int(max_row) + extra_rows
        ws.merge_cells(
            f"{min_col}{min_row}:{max_col}{max_row}"
        )


def prepare_template(ws, record_count: int) -> int:
    template_first_data_row = 7
    template_last_data_row = 36
    total_row_template = 37

    if record_count > template_last_data_row - template_first_data_row + 1:
        extra_rows = record_count - (template_last_data_row - template_first_data_row + 1)
        shift_footer_merges_for_extra_rows(ws, extra_rows)
    else:
        extra_rows = 0

    data_last_row = template_first_data_row + record_count - 1 if record_count else 6
    total_row = total_row_template + extra_rows

    # Clear template sample values without disturbing formatting or merged cells.
    clear_last = max(template_last_data_row + extra_rows, data_last_row)
    for row in range(template_first_data_row, clear_last + 1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row, col)
            if isinstance(cell, MergedCell):
                continue
            cell.value = None

    # Add styles/merges for rows beyond the original template data area.
    if extra_rows:
        for row in range(template_last_data_row + 1, template_last_data_row + extra_rows + 1):
            copy_row_format(ws, template_last_data_row, row)
            ws.merge_cells(f"B{row}:D{row}")

    # Ensure every output data row has the same basic merge as the template.
    for row in range(template_first_data_row, data_last_row + 1):
        merge_ref = f"B{row}:D{row}"
        if merge_ref not in {str(x) for x in ws.merged_cells.ranges}:
            ws.merge_cells(merge_ref)

    return total_row


def write_workbook(
    records: list[dict[str, Any]],
    service_charge: float = 500,
    cc_tada_per_pdf: float = 100,
    output_name: str = "Generated_Annexure_XXI.xlsx",
) -> io.BytesIO:
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"Template not found: {TEMPLATE}")

    wb = load_workbook(TEMPLATE)
    ws = wb["Vendor Claim Sheet"]

    total_row = prepare_template(ws, len(records))
    start_row = 7
    seen_sources: set[str] = set()

    for idx, record in enumerate(records, start=1):
        row = start_row + idx - 1
        source_pdf = str(record.get("_source_pdf") or "")

        ws.cell(row, 1).value = idx
        ws.cell(row, 2).value = record.get("machine_number") or ""
        ws.cell(row, 5).value = record.get("capacity") or ""
        ws.cell(row, 6).value = record.get("class") or ""
        ws.cell(row, 7).value = record.get("customer_name") or ""
        ws.cell(row, 8).value = record.get("location") or "HOSUR"
        ws.cell(row, 9).value = record.get("cv_number") or ""

        parsed_date = parse_date(record.get("cv_date"))
        ws.cell(row, 10).value = parsed_date
        if isinstance(parsed_date, datetime):
            ws.cell(row, 10).number_format = "dd-mm-yyyy"

        actual = record.get("actual_stamping_fee")
        site = record.get("site_stamping_fee")
        od = record.get("od_charges")

        ws.cell(row, 11).value = actual
        ws.cell(row, 12).value = site if site not in ("", None, 0, 0.0) else None

        # ₹100 CC / TA-DA is charged once for each uploaded PDF.
        cc_value = cc_tada_per_pdf if source_pdf not in seen_sources else None
        ws.cell(row, 13).value = cc_value
        seen_sources.add(source_pdf)

        ws.cell(row, 14).value = od if od not in ("", None, 0, 0.0) else None
        ws.cell(row, 15).value = None

        ws.cell(row, 16).value = f"=K{row}+L{row}+M{row}+N{row}+O{row}"
        ws.cell(row, 17).value = service_charge
        ws.cell(row, 18).value = f"=P{row}+Q{row}"

    end_data_row = 6 + len(records)
    if end_data_row >= start_row:
        end_ref = end_data_row
        for col, formula_col in zip(range(11, 19), "KLMNOPQR"):
            ws.cell(total_row, col).value = f"=SUM({formula_col}{start_row}:{formula_col}{end_ref})"
    else:
        for col in range(11, 19):
            ws.cell(total_row, col).value = 0

    # Encourage spreadsheet applications to recalculate formulas on open.
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def decode_records(records_json: str | None) -> list[dict[str, Any]] | None:
    if not records_json:
        return None
    try:
        payload = json.loads(records_json)
        if not isinstance(payload, list):
            raise ValueError
        return [dict(item) for item in payload]
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Invalid records_json payload.") from exc


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok", "version": "2.0.0"}


@app.post("/extract")
async def extract(
    request: Request,
    files: list[UploadFile] = File(...),
):
    check_password(request)

    if not files:
        raise HTTPException(status_code=400, detail="Upload at least one PDF.")

    results = []
    for file in files:
        data = await file.read()
        if not (file.filename or "").lower().endswith(".pdf"):
            raise HTTPException(status_code=400, detail=f"{file.filename}: only PDF files are supported.")
        results.append(extract_file(data, file.filename or "document.pdf"))

    return {"files": results}


@app.post("/generate")
async def generate(
    request: Request,
    files: list[UploadFile] = File(...),
    service_charge: float = Form(500),
    cc_tada_per_pdf: float = Form(100),
    output_mode: str = Form("combined"),
    records_json: str | None = Form(None),
):
    check_password(request)

    if service_charge < 0 or cc_tada_per_pdf < 0:
        raise HTTPException(status_code=400, detail="Charges cannot be negative.")

    if output_mode not in {"combined", "separate"}:
        raise HTTPException(status_code=400, detail="output_mode must be combined or separate.")

    decoded = decode_records(records_json)
    if decoded is not None:
        if not files:
            raise HTTPException(status_code=400, detail="Source PDFs are required.")
        records = decoded
        # Validate source PDF names against the upload set.
        source_names = {f.filename or "" for f in files}
        for rec in records:
            if rec.get("_source_pdf") not in source_names:
                raise HTTPException(
                    status_code=400,
                    detail=f"Record refers to an unknown source PDF: {rec.get('_source_pdf')}",
                )
    else:
        file_groups = []
        for file in files:
            data = await file.read()
            file_groups.append(extract_file(data, file.filename or "document.pdf"))
        records = [
            certificate
            for group in file_groups
            for certificate in group["certificates"]
        ]

    if not records:
        raise HTTPException(
            status_code=422,
            detail="No certificates were detected in the uploaded PDF(s).",
        )

    if output_mode == "combined":
        output = write_workbook(
            records,
            service_charge=service_charge,
            cc_tada_per_pdf=cc_tada_per_pdf,
        )
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": 'attachment; filename="Generated_Annexure_XXI.xlsx"'
            },
        )

    # Separate Excel workbook for each uploaded PDF.
    grouped: dict[str, list[dict[str, Any]]] = {}
    for record in records:
        grouped.setdefault(
            str(record.get("_source_pdf") or "document.pdf"),
            [],
        ).append(record)

    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for source_name, group in grouped.items():
            workbook = write_workbook(
                group,
                service_charge=service_charge,
                cc_tada_per_pdf=cc_tada_per_pdf,
                output_name=f"{filename_stem(source_name)}_Annexure_XXI.xlsx",
            )
            zf.writestr(
                f"{filename_stem(source_name)}_Annexure_XXI.xlsx",
                workbook.getvalue(),
            )

    archive.seek(0)
    return StreamingResponse(
        archive,
        media_type="application/zip",
        headers={
            "Content-Disposition": 'attachment; filename="Generated_Annexure_XXI_by_PDF.zip"'
        },
    )
